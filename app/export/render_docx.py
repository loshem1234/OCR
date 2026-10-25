"""Word (.docx) renderer."""
from __future__ import annotations

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..db import Document
from ..ocr.textutil import is_rtl
from .assemble import Assembled, Block
from .inline import tokens
from .options import ExportOptions

FONT = "Times New Roman"  # ships with Word on Windows and macOS; covers polytonic Greek, Hebrew, Arabic
GOLD = RGBColor(0x8A, 0x7B, 0x55)
GREY = RGBColor(0x6B, 0x6B, 0x6B)


def _set_fonts(run_or_style, name: str) -> None:
    run_or_style.font.name = name
    rpr = run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _rtl(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _add_inline(paragraph, text: str, opts: ExportOptions, rtl: bool) -> None:
    for kind, v in tokens(text, opts.markup):
        if kind == "br":
            paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        if kind == "page":
            v = opts.fmt_marker(v)
        run = paragraph.add_run(v)
        if rtl:
            rpr = run._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:rtl"))
        if kind == "italic":
            run.italic = True
        elif kind == "bold":
            run.bold = True
        elif kind == "fn":
            run.font.superscript = True
        elif kind == "ref":
            run.font.superscript = True
            run.font.color.rgb = GREY
        elif kind == "page":
            run.font.size = Pt(max(opts.font_size - 3, 7))
            run.font.color.rgb = GOLD


def _para(container, text: str, opts: ExportOptions, style: str | None = None):
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    rtl = is_rtl(text)
    if rtl:
        _rtl(p)
    elif "\n" not in text and style is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_inline(p, text, opts, rtl)
    return p


def _blocks(container, blocks: list[Block], opts: ExportOptions, doc) -> None:
    for b in blocks:
        if b.kind == "heading":
            if container is doc:
                h = doc.add_heading(level=min(b.level, 4))
                _add_inline(h, b.text, opts, is_rtl(b.text))
            else:
                p = _para(container, b.text, opts)
                for r in p.runs:
                    r.bold = True
        elif b.kind in ("para", "marker"):
            _para(container, b.text, opts)
        elif b.kind == "facing":
            table = container.add_table(rows=1, cols=2)
            left, right = table.rows[0].cells
            for cell, side in ((left, b.left), (right, b.right)):
                cell.paragraphs[0].text = ""
                _blocks(cell, side, opts, doc)
                if len(cell.paragraphs) > 1 and not cell.paragraphs[0].text:
                    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        elif b.kind == "notes":
            lab = container.add_paragraph()
            r = lab.add_run("Notes " + opts.fmt_marker(b.label))
            r.font.size, r.font.color.rgb, r.italic = Pt(opts.font_size - 2), GOLD, True
            for t in b.text.split("\n\n"):
                if t.strip():
                    p = _para(container, t, opts)
                    for run in p.runs:
                        run.font.size = Pt(opts.font_size - 1.5)


def render_docx(a: Assembled, doc: Document, opts: ExportOptions, path: str) -> None:
    d = Docx()
    normal = d.styles["Normal"]
    _set_fonts(normal, FONT)
    normal.font.size = Pt(opts.font_size)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            st = d.styles[name]
            _set_fonts(st, FONT)
            st.font.color.rgb = RGBColor(0x22, 0x1E, 0x18)
        except KeyError:
            pass
    d.core_properties.title = doc.title
    d.core_properties.author = doc.author or ""

    if opts.title_page:
        d.add_heading(doc.title, level=0)
        if doc.author:
            p = d.add_paragraph()
            p.add_run(doc.author).italic = True
        d.add_page_break()

    _blocks(d, a.blocks, opts, d)

    if a.notes:
        d.add_page_break()
        d.add_heading("Notes", level=1)
        for label, text in a.notes:
            lab = d.add_paragraph()
            r = lab.add_run(opts.fmt_marker(label))
            r.bold, r.font.color.rgb = True, GOLD
            for t in text.split("\n\n"):
                if t.strip():
                    _para(d, t, opts)
    d.save(path)
